"""Source material for an episode: text out of .txt / .md / .pdf / .docx uploads, and the chunk-and-summarise
pass that keeps long material inside what a local LLM can read in one prompt."""
import io
import re
import zipfile

MB = 1024 * 1024
MAX_BYTES = 20 * MB
MAX_CHARS = 400_000
MAX_PAGES = 500
# above this the writer reads summaries instead of the material itself
DIRECT_CHARS = 12_000
CHUNK_CHARS = 6_000
MAX_CHUNKS = 40

TEXT_EXT = {".txt": "txt", ".md": "md", ".markdown": "md", ".text": "txt"}


class IngestError(ValueError):
    """Readable; the API turns it into a 422."""


def _ext(name: str) -> str:
    m = re.search(r"\.[A-Za-z0-9]+$", name or "")
    return m[0].lower() if m else ""


def kind_for(name: str, head: bytes) -> str:
    ext = _ext(name)
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04") and ext in (".docx", ""):
        return "docx"
    if ext in TEXT_EXT:
        return TEXT_EXT[ext]
    if ext == ".doc":
        raise IngestError("Old .doc files can't be read; save it as .docx or PDF")
    raise IngestError("Upload a .txt, .md, .pdf or .docx file")


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    if b"\x00" in data[:4096]:
        raise IngestError("That doesn't look like a text file")
    return data.decode("cp1252", errors="replace")


def _pdf(data: bytes) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError:
        raise IngestError("PDF support isn't installed on the server (pypdf)") from None
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise IngestError("That PDF is password protected")
        pages = reader.pages
        n = len(pages)
        if n > MAX_PAGES:
            raise IngestError(f"That PDF has {n} pages; {MAX_PAGES} is the most")
        parts = [(p.extract_text() or "").strip() for p in pages]
    except (PdfReadError, ValueError, KeyError, OSError) as e:
        raise IngestError(f"That PDF couldn't be read ({e})") from None
    text = "\n\n".join(p for p in parts if p)
    if not text.strip():
        raise IngestError("That PDF has no text layer (is it a scan?)")
    return text, n


def _docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise IngestError("Word support isn't installed on the server (python-docx)") from None
    try:
        doc = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as e:
        raise IngestError(f"That Word file couldn't be read ({e})") from None
    parts = [p.text.strip() for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(p for p in parts if p)
    if not text.strip():
        raise IngestError("That Word file has no text in it")
    return text


def tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract(name: str, data: bytes) -> dict:
    """{type, text, chars, pages?, truncated} for one uploaded file."""
    if not data:
        raise IngestError("That file is empty")
    if len(data) > MAX_BYTES:
        raise IngestError(f"Documents can be up to {MAX_BYTES // MB} MB")
    kind = kind_for(name, data[:8])
    pages = None
    if kind == "pdf":
        text, pages = _pdf(data)
    elif kind == "docx":
        text = _docx(data)
    else:
        text = _decode(data)
    text = tidy(text)
    if not text:
        raise IngestError("There's no text in that file")
    out = {"type": kind, "text": text[:MAX_CHARS], "chars": min(len(text), MAX_CHARS),
           "truncated": len(text) > MAX_CHARS}
    if pages is not None:
        out["pages"] = pages
    return out


def material(sources: dict) -> str:
    """The pasted text and every document, each under its own heading."""
    parts = []
    if (sources.get("text") or "").strip():
        parts.append(sources["text"].strip())
    for d in sources.get("docs") or []:
        if (d.get("text") or "").strip():
            parts.append(f"## {d.get('name') or 'Document'}\n\n{d['text'].strip()}")
    return "\n\n".join(parts)[:MAX_CHARS]


def chunks(text: str, size: int = CHUNK_CHARS) -> list[str]:
    """Pieces of about `size` characters, cut at paragraph (else sentence) ends."""
    out, cur = [], ""
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        while len(para) > size:
            cut = max(para.rfind(". ", 0, size), para.rfind("۔", 0, size), para.rfind("\n", 0, size))
            cut = cut + 1 if cut > size // 3 else size
            if cur:
                out.append(cur)
                cur = ""
            out.append(para[:cut].strip())
            para = para[cut:].strip()
        if cur and len(cur) + len(para) + 2 > size:
            out.append(cur)
            cur = ""
        cur = f"{cur}\n\n{para}" if cur else para
    if cur:
        out.append(cur)
    return out[:MAX_CHUNKS]


def summary_messages(chunk: str, n: int, total: int, topic: str) -> list[dict]:
    rules = [
        "You condense source material for a podcast writer.",
        "List the facts, names, numbers, dates, quotes and arguments in this part as short plain bullet points. "
        "Keep every name and number exactly as written. Add nothing that isn't in the text.",
        "Write the bullets in the language of the text.",
    ]
    head = f"Part {n} of {total}" + (f" (the episode is about: {topic})" if topic else "")
    return [{"role": "system", "content": "\n".join(rules)}, {"role": "user", "content": f"{head}\n\n{chunk}"}]


def condensed(text: str, topic: str, tick=None) -> tuple[str, int]:
    """The material itself when it's short, else the joined summaries of its chunks. (text, chunks summarised)"""
    if len(text) <= DIRECT_CHARS:
        return text, 0
    from app.llm import chat_sync

    parts = chunks(text)
    notes = []
    for n, part in enumerate(parts, start=1):
        if tick:
            tick(n, len(parts))
        res = chat_sync("creative", summary_messages(part, n, len(parts), topic), think=False, temperature=0.2,
                        max_tokens=900, tick=(lambda: tick(n, len(parts))) if tick else None)
        notes.append(res.content.strip())
    return "\n\n".join(notes), len(parts)
